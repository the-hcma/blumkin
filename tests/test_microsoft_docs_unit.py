"""Hermetic tests for the Microsoft `docs create` backend and its scope gate."""

from __future__ import annotations

import asyncio
import io
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document
from kiota_abstractions.method import Method
from msgraph.generated.models.o_data_errors.o_data_error import ODataError

from blumkin.config import BlumkinConfig, MailSignatureConfig
from blumkin.created_docs import is_blumkin_created_doc, record_created_doc
from blumkin.providers import microsoft_docs as _md
from blumkin.providers.kind import ProviderKind
from blumkin.providers.microsoft import MicrosoftWorkspaceProvider
from blumkin.skills.dispatch import run_skill
from blumkin.skills.docs import DocBodyError
from blumkin.skills.drive import DriveItemNotFoundError
from blumkin.skills.errors import ScopeAddonDisabledError

_MOD = "blumkin.providers.microsoft_docs"
_UPLOADED = SimpleNamespace(
    id="01ABC", name="Brief.docx", web_url="https://onedrive.example/Brief.docx"
)


@pytest.fixture(autouse=True)
def _isolate_created_docs(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep the `docs create` ownership record out of the repo (default cfg dir is fake)."""
    monkeypatch.setattr(
        BlumkinConfig,
        "created_docs_path",
        property(lambda _self: tmp_path / "created_docs.json"),
    )


def _cfg(
    *,
    provider: ProviderKind = ProviderKind.MICROSOFT,
    docs_scopes: bool = True,
    config_dir: Path = Path("unused"),
) -> BlumkinConfig:
    return BlumkinConfig(
        client_id="abc",
        config_dir=config_dir,
        default_tz="UTC",
        docs_scopes=docs_scopes,
        email="",
        files_scopes=False,
        google_oauth_client_file=None,
        graph_timeout_seconds=60.0,
        legacy_flat=True,
        mail_signature=MailSignatureConfig(),
        profile="default",
        provider=provider,
        tags=(),
        tenant_id="brk.tech",
        wo1162425_scopes=False,
    )


def _client(*, folder_exists: bool = False) -> MagicMock:
    client = MagicMock()
    calls: list[Any] = []

    async def send_async(request_info: Any, _factory: Any, _error_map: Any) -> Any:
        calls.append(request_info)
        if ":/content" in request_info.url_template:
            return _UPLOADED
        if request_info.http_method == Method.GET:  # folder existence probe
            return SimpleNamespace(id="existing-folder") if folder_exists else None
        return SimpleNamespace(id="new-folder")  # POST children (folder create)

    client.request_adapter.send_async = AsyncMock(side_effect=send_async)
    client.calls = calls
    return client


def _run(client: MagicMock, monkeypatch: pytest.MonkeyPatch, **kwargs: Any) -> dict[str, Any]:
    monkeypatch.setattr(f"{_MOD}.create_graph_client", lambda _cfg: client)
    return asyncio.run(MicrosoftWorkspaceProvider(_cfg()).docs_create(**kwargs))


def _style_names(document: Any) -> list[str]:
    return [p.style.name if p.style else "" for p in document.paragraphs]


def test_docs_create_creates_a_missing_folder(monkeypatch: pytest.MonkeyPatch) -> None:
    client = _client(folder_exists=False)
    payload = _run(client, monkeypatch, title="Weekly: status?", body="hi", folder="/Reports/Q3/")
    assert payload["document"]["folder"] == "Reports/Q3"
    gets = [c for c in client.calls if c.http_method == Method.GET]
    posts = [c for c in client.calls if c.http_method == Method.POST]
    assert [g.path_parameters["path"] for g in gets] == ["Reports", "Reports/Q3"]
    assert len(posts) == 2  # both segments created
    (upload,) = [c for c in client.calls if ":/content" in c.url_template]
    assert upload.path_parameters["path"] == "Reports/Q3/Weekly status.docx"


def test_docs_create_on_google_bypasses_the_docs_scopes_gate() -> None:
    provider = MagicMock()
    provider.docs_create = AsyncMock(return_value={"document": {"id": "d1"}})
    payload = asyncio.run(
        run_skill(
            "docs.create",
            {"title": "T", "body": "x"},
            config=_cfg(provider=ProviderKind.GOOGLE, docs_scopes=False),
            provider=provider,
        )
    )
    assert payload == {"document": {"id": "d1"}}
    provider.docs_create.assert_awaited_once()


def test_docs_create_on_microsoft_needs_the_docs_scopes_toggle() -> None:
    with pytest.raises(ScopeAddonDisabledError, match="Files.ReadWrite"):
        asyncio.run(
            run_skill(
                "docs.create",
                {"title": "T", "body": "x"},
                config=_cfg(docs_scopes=False),
                provider=MicrosoftWorkspaceProvider(_cfg(docs_scopes=False)),
            )
        )


def test_docs_create_propagates_a_non_404_folder_probe_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client = MagicMock()
    throttled = ODataError()
    throttled.response_status_code = 429

    async def send_async(request_info: Any, _f: Any, _e: Any) -> Any:
        if request_info.http_method == Method.GET:
            raise throttled
        raise AssertionError("must not write after a throttled probe")

    client.request_adapter.send_async = AsyncMock(side_effect=send_async)
    monkeypatch.setattr(f"{_MOD}.create_graph_client", lambda _cfg: client)
    with pytest.raises(ODataError):
        asyncio.run(
            MicrosoftWorkspaceProvider(_cfg()).docs_create(title="T", body="hi", folder="Reports")
        )


def test_docs_create_reads_a_body_file_as_plain_text(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    src = tmp_path / "note.txt"
    src.write_text("# not a heading\n\nsecond para", encoding="utf-8")
    client = _client()
    _run(client, monkeypatch, title="T", body_file=str(src), body_format="text")
    (upload,) = [c for c in client.calls if ":/content" in c.url_template]
    texts = [p.text for p in Document(io.BytesIO(bytes(upload.content))).paragraphs]
    assert texts == ["# not a heading", "second para"]


def test_docs_create_renders_every_supported_block(monkeypatch: pytest.MonkeyPatch) -> None:
    body = (
        "# Title\n\n"
        "para text\n\n"
        "- b0\n  - b1\n\n"
        "1. n0\n\n"
        "---\n\n"
        "```\ncode line\n```\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n"
    )
    client = _client()
    _run(client, monkeypatch, title="T", body=body)
    (upload,) = [c for c in client.calls if ":/content" in c.url_template]
    document = Document(io.BytesIO(bytes(upload.content)))
    styles = _style_names(document)
    assert "Heading 1" in styles
    assert "List Bullet" in styles and "List Bullet 2" in styles
    assert "1. n0" in [p.text for p in document.paragraphs]  # numbered marker is literal text


def test_docs_create_restarts_numbering_for_each_ordered_list(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client = _client()
    _run(client, monkeypatch, title="T", body="1. a\n2. b\n\n# Next\n\n1. c\n2. d")
    (upload,) = [c for c in client.calls if ":/content" in c.url_template]
    texts = [p.text for p in Document(io.BytesIO(bytes(upload.content))).paragraphs]
    # The second list restarts at 1, it does not continue as 3./4.
    assert texts == ["1. a", "2. b", "Next", "1. c", "2. d"]


def test_docs_create_renders_markdown_links_as_real_hyperlinks(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client = _client()
    _run(client, monkeypatch, title="T", body="See [the guide](https://x.test/guide) now.")
    (upload,) = [c for c in client.calls if ":/content" in c.url_template]
    (paragraph,) = Document(io.BytesIO(bytes(upload.content))).paragraphs
    assert paragraph.text == "See the guide now."
    # A real w:hyperlink with an external relationship, not just underlined text.
    assert [(h.address, h.text) for h in paragraph.hyperlinks] == [
        ("https://x.test/guide", "the guide")
    ]


def test_docs_create_reuses_an_existing_folder(monkeypatch: pytest.MonkeyPatch) -> None:
    client = _client(folder_exists=True)
    _run(client, monkeypatch, title="T", body="hi", folder="Reports")
    assert [c for c in client.calls if c.http_method == Method.POST] == []


def test_docs_create_strips_url_breaking_chars_from_title_and_folder(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client = _client(folder_exists=True)
    payload = _run(client, monkeypatch, title="Sprint #12: plan?", body="hi", folder="Team #2/Q4")
    assert payload["document"]["folder"] == "Team 2/Q4"
    (upload,) = [c for c in client.calls if ":/content" in c.url_template]
    assert upload.path_parameters["path"] == "Team 2/Q4/Sprint 12 plan.docx"
    assert "#" not in upload.path_parameters["path"]


def test_docs_create_uploads_a_real_docx(monkeypatch: pytest.MonkeyPatch) -> None:
    client = _client()
    payload = _run(
        client, monkeypatch, title="Brief", body="# Brief\n\n- one\n- two\n\n`code` and **bold**"
    )
    assert payload["document"] == {
        "id": "01ABC",
        "name": "Brief.docx",
        "web_url": "https://onedrive.example/Brief.docx",
        "provider": "microsoft",
        "format": "docx",
        "folder": None,
    }
    (upload,) = [c for c in client.calls if ":/content" in c.url_template]
    assert upload.path_parameters["path"] == "Brief.docx"
    assert upload.http_method == Method.PUT
    # Never silently overwrite an existing file at the same title.
    assert "conflictBehavior=rename" in upload.url_template
    document = Document(io.BytesIO(bytes(upload.content)))
    assert document.paragraphs[0].text == "Brief"
    assert _style_names(document)[0] == "Heading 1"


# --------------------------------------------------------------------------- docs update


def _update_client(*, missing: bool = False, patch_name: str = "Renamed.docx") -> MagicMock:
    client = MagicMock()
    calls: list[Any] = []

    async def send_async(request_info: Any, _factory: Any, _error_map: Any) -> Any:
        calls.append(request_info)
        if request_info.http_method == Method.GET:
            if missing:
                err = ODataError()
                err.response_status_code = 404
                raise err
            return SimpleNamespace(id="01ABC", name="Brief.docx", web_url="https://od/Brief.docx")
        if request_info.http_method == Method.PATCH:
            return SimpleNamespace(id="01ABC", name=patch_name, web_url="https://od/renamed")
        return SimpleNamespace(id="01ABC", name="Brief.docx", web_url="https://od/Brief.docx")

    client.request_adapter.send_async = AsyncMock(side_effect=send_async)
    client.calls = calls
    return client


def _cfg_with_doc(doc_id: str = "01ABC") -> BlumkinConfig:
    cfg = _cfg()
    record_created_doc(cfg, doc_id)
    return cfg


def _run_update(
    client: MagicMock,
    monkeypatch: pytest.MonkeyPatch,
    *,
    rename_spy: list[tuple[str, str]] | None = None,
    **kwargs: Any,
) -> dict[str, Any]:
    monkeypatch.setattr(f"{_MOD}.create_graph_client", lambda _cfg: client)
    if rename_spy is not None:
        real = _md._rename_item

        async def spy(c: Any, item_id: str, name: str) -> Any:
            rename_spy.append((item_id, name))
            return await real(c, item_id, name)

        monkeypatch.setattr(_md, "_rename_item", spy)
    return asyncio.run(MicrosoftWorkspaceProvider(_cfg_with_doc()).docs_update(**kwargs))


def test_docs_update_replaces_bytes_via_put_content(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    client = _update_client()
    payload = _run_update(client, monkeypatch, document_id="01ABC", body="# New\n\nbody")
    methods = [c.http_method for c in client.calls]
    assert methods == [Method.GET, Method.PUT]
    (put,) = [c for c in client.calls if c.http_method == Method.PUT]
    assert put.url_template.endswith("/items/{id}/content")
    assert put.path_parameters["id"] == "01ABC"
    assert Document(io.BytesIO(bytes(put.content))).paragraphs[0].text == "New"
    assert payload["document"] == {
        "id": "01ABC",
        "name": "Brief.docx",
        "web_url": "https://od/Brief.docx",
        "provider": "microsoft",
        "format": "docx",
        "folder": None,
    }


def test_docs_update_renames_via_patch_with_a_sanitized_name(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    client = _update_client(patch_name="Q3 plan.docx")
    spy: list[tuple[str, str]] = []
    payload = _run_update(
        client, monkeypatch, rename_spy=spy, document_id="01ABC", title="Q3: plan?"
    )
    assert [c.http_method for c in client.calls] == [Method.GET, Method.PATCH]
    # The real production name string, not the mock's canned return value.
    assert spy == [("01ABC", "Q3 plan.docx")]
    assert payload["document"]["name"] == "Q3 plan.docx"


def test_docs_update_title_and_body_together(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    client = _update_client()
    _run_update(client, monkeypatch, document_id="01ABC", title="Q3", body="hi")
    assert [c.http_method for c in client.calls] == [Method.GET, Method.PUT, Method.PATCH]


def test_docs_update_does_not_double_a_docx_suffix_in_the_title(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    spy: list[tuple[str, str]] = []
    _run_update(_update_client(), monkeypatch, rename_spy=spy, document_id="01ABC", title="Q3.docx")
    assert spy == [("01ABC", "Q3.docx")]  # not "Q3.docx.docx"


def test_docs_update_reads_a_body_file(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    src = tmp_path / "body.md"
    src.write_text("# From a file\n\nfile body text", encoding="utf-8")
    client = _update_client()
    _run_update(client, monkeypatch, document_id="01ABC", body_file=str(src))
    (put,) = [c for c in client.calls if c.http_method == Method.PUT]
    texts = [p.text for p in Document(io.BytesIO(bytes(put.content))).paragraphs]
    assert "From a file" in texts and "file body text" in texts


def test_docs_update_needs_a_change(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    with pytest.raises(DocBodyError, match="at least one of"):
        _run_update(_update_client(), monkeypatch, document_id="01ABC")


def test_docs_update_id_deleted_since_creation_is_not_found(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    with pytest.raises(DriveItemNotFoundError):
        _run_update(_update_client(missing=True), monkeypatch, document_id="01ABC", body="hi")


def test_docs_update_refuses_a_doc_this_install_did_not_create(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    client = _update_client()
    monkeypatch.setattr(f"{_MOD}.create_graph_client", lambda _cfg: client)
    with pytest.raises(DriveItemNotFoundError, match="created by this blumkin"):
        asyncio.run(
            MicrosoftWorkspaceProvider(_cfg()).docs_update(document_id="0xFOREIGN", body="clobber")
        )
    assert client.calls == []  # bailed before any Graph call - nothing overwritten


def test_docs_update_refuses_to_rewrite_a_non_docx_item(monkeypatch: pytest.MonkeyPatch) -> None:
    """Even a recorded id must resolve to a .docx before its bytes are replaced."""
    client = _update_client()

    async def send_async(request_info: Any, _f: Any, _e: Any) -> Any:
        client.calls.append(request_info)
        if request_info.http_method == Method.GET:
            return SimpleNamespace(id="01ABC", name="notes.txt", web_url="https://od/notes.txt")
        raise AssertionError("must not write to a non-docx item")

    client.request_adapter.send_async = AsyncMock(side_effect=send_async)
    with pytest.raises(DriveItemNotFoundError, match="not a .docx"):
        _run_update(client, monkeypatch, document_id="01ABC", body="clobber")


def test_docs_create_records_the_new_id(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    client = _client()
    monkeypatch.setattr(f"{_MOD}.create_graph_client", lambda _cfg: client)
    asyncio.run(MicrosoftWorkspaceProvider(_cfg()).docs_create(title="T", body="x"))
    assert is_blumkin_created_doc(_cfg(), "01ABC")


def test_docs_update_on_microsoft_needs_the_docs_scopes_toggle() -> None:
    with pytest.raises(ScopeAddonDisabledError, match="Files.ReadWrite"):
        asyncio.run(
            run_skill(
                "docs.update",
                {"id": "01ABC", "body": "x"},
                config=_cfg(docs_scopes=False),
                provider=MicrosoftWorkspaceProvider(_cfg(docs_scopes=False)),
            )
        )
