"""Microsoft `docs create`: build a .docx from the shared block model, upload to OneDrive.

Graph has no Word-authoring API, so the ``.docx`` (OOXML) is built in memory with
``python-docx`` from the same :class:`~blumkin.skills.docs.DocBlock` list the
Google backend renders, then uploaded with a simple
``PUT /me/drive/root:/{path}:/content``. Needs ``Files.ReadWrite`` (the
``docs_scopes`` opt-in); the dispatch layer refuses the skill before this runs
when the toggle is off.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from kiota_abstractions.method import Method
from kiota_abstractions.request_information import RequestInformation
from kiota_abstractions.serialization.parsable_factory import ParsableFactory
from msgraph.generated.models.drive_item import DriveItem
from msgraph.generated.models.folder import Folder
from msgraph.generated.models.o_data_errors.o_data_error import ODataError

from blumkin.config import BlumkinConfig, load_config
from blumkin.created_docs import is_blumkin_created_doc, record_created_doc
from blumkin.graph import create_graph_client
from blumkin.skills.docs import (
    DocBlock,
    DocSpan,
    parse_body,
    read_body,
    read_update_body,
    require_docs_update_target,
    strip_docx_suffix,
    table_to_text,
)
from blumkin.skills.drive import DriveItemNotFoundError


async def docs_create(
    *,
    title: str,
    body: str | None = None,
    body_file: str | None = None,
    body_format: str = "markdown",
    folder: str | None = None,
    config: BlumkinConfig | None = None,
) -> dict[str, Any]:
    """Build ``title``.docx from authored ``body`` and upload it to the user's drive."""
    if not title.strip():
        raise ValueError("--title is required")
    cfg = config or load_config()
    blocks = parse_body(read_body(body, body_file), body_format=body_format)

    name = _docx_name(title)
    folder_name = _safe_folder(folder)
    item_path = f"{folder_name}/{name}" if folder_name else name

    client = create_graph_client(cfg)
    if folder_name:
        # Graph's simple upload does not auto-create missing parents, and the
        # catalog documents --folder as "created if absent" (parity with Google).
        await _ensure_folder(client, folder_name)
    request_info = RequestInformation(Method.PUT, _UPLOAD_URL, {"path": item_path})
    request_info.set_stream_content(_render_docx(blocks), _DOCX_MIME)
    item = await client.request_adapter.send_async(request_info, DriveItem, _ERROR_MAP)
    if item is None:
        raise RuntimeError("Graph returned no DriveItem for the uploaded document")

    record_created_doc(cfg, item.id or "")
    return {
        "document": {
            "id": item.id,
            "name": item.name or name,
            "web_url": item.web_url,
            "provider": "microsoft",
            "format": "docx",
            "folder": folder_name,
        }
    }


async def docs_update(
    *,
    document_id: str,
    title: str | None = None,
    body: str | None = None,
    body_file: str | None = None,
    body_format: str = "markdown",
    config: BlumkinConfig | None = None,
) -> dict[str, Any]:
    """Re-render a .docx blumkin uploaded: rename the DriveItem, replace its bytes, or both.

    Graph has no partial Word update, so a body change rebuilds the whole ``.docx``
    from the new content (``python-docx``, same as ``docs create``) and
    ``PUT``s it to the same DriveItem - id, URL, and permissions are unchanged.

    Refuses an id this blumkin install did not upload: ``Files.ReadWrite`` covers
    the whole drive, so without the local ``docs create`` record ``--id`` alone
    could point at (and overwrite the bytes of) any file the account owns. A
    document created on another machine is not updatable here - open it in the
    browser.
    """
    doc_id = document_id.strip()
    if not doc_id:
        raise ValueError("--id is required")
    new_title = require_docs_update_target(title=title, body=body, body_file=body_file)
    raw_body = read_update_body(body, body_file)
    blocks = parse_body(raw_body, body_format=body_format) if raw_body is not None else None

    cfg = config or load_config()
    if not is_blumkin_created_doc(cfg, doc_id):
        raise DriveItemNotFoundError(
            f"no document with id {doc_id!r} was created by this blumkin - `docs update` "
            "only touches documents this install uploaded (open others in the browser)"
        )
    client = create_graph_client(cfg)
    item = await _get_item_by_id(client, doc_id)
    # Defence in depth against a stale / hand-edited record: only ever rewrite a
    # .docx (what `docs create` uploads), never some other file at that id.
    if blocks is not None and not (item.name or "").lower().endswith(".docx"):
        raise DriveItemNotFoundError(
            f"item {doc_id!r} is not a .docx uploaded by `docs create` ({item.name!r})"
        )

    if blocks is not None:
        put_info = RequestInformation(Method.PUT, _CONTENT_BY_ID_URL, {"id": doc_id})
        put_info.set_stream_content(_render_docx(blocks), _DOCX_MIME)
        item = await client.request_adapter.send_async(put_info, DriveItem, _ERROR_MAP) or item

    if new_title is not None:
        item = await _rename_item(client, doc_id, _docx_name(new_title)) or item

    return {
        "document": {
            "id": doc_id,
            "name": item.name,
            "web_url": item.web_url,
            "provider": "microsoft",
            "format": "docx",
            "folder": None,
        }
    }


_CHILDREN_ROOT_URL = "https://graph.microsoft.com/v1.0/me/drive/root/children"
_CHILDREN_URL = "https://graph.microsoft.com/v1.0/me/drive/root:/{+path}:/children"
_CODE_FONT = "Consolas"
_CONTENT_BY_ID_URL = "https://graph.microsoft.com/v1.0/me/drive/items/{id}/content"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_ERROR_MAP: dict[str, type[ParsableFactory]] = {"4XX": ODataError, "5XX": ODataError}
_HYPERLINK_COLOR = "0563C1"
_ITEM_BY_ID_URL = "https://graph.microsoft.com/v1.0/me/drive/items/{id}"
_ITEM_URL = "https://graph.microsoft.com/v1.0/me/drive/root:/{+path}"
# Bullet lists keep the built-in Word style (no shared counter to fight); numbered
# lists render their marker as literal text - see `_render_block`.
_LIST_STYLES = {0: "List Bullet", 1: "List Bullet 2"}
# The upload URL uses RFC 6570 reserved expansion ({+path}), which leaves `#`,
# `?`, `:` etc. literal - a raw `#` in a title or folder truncates the request
# path at the fragment. Strip both the OneDrive-illegal set and these.
_UNSAFE_SEGMENT_CHARS = frozenset('\\/:*?"<>|#%')
# conflictBehavior=rename: never silently overwrite an existing file at the same
# path - a repeated `docs create --title X` produces "X 1.docx", not a clobber
# (closer to Google, where documents.create always mints a new doc).
_UPLOAD_URL = (
    "https://graph.microsoft.com/v1.0/me/drive/root:/{+path}:/content"
    "?%40microsoft.graph.conflictBehavior=rename"
)


def _add_hyperlink(paragraph: Paragraph, *, url: str, text: str, bold: bool, italic: bool) -> None:
    """Append ``text`` as a real external-relationship w:hyperlink (clickable)."""
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _HYPERLINK_COLOR)
    run_properties.append(color)
    if bold:
        run_properties.append(OxmlElement("w:b"))
    if italic:
        run_properties.append(OxmlElement("w:i"))
    run.append(run_properties)

    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run.append(text_element)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_runs(paragraph: Paragraph, spans: tuple[DocSpan, ...]) -> None:
    for span in spans:
        if span.link:
            _add_hyperlink(
                paragraph, url=span.link, text=span.text, bold=span.bold, italic=span.italic
            )
            continue
        run = paragraph.add_run(span.text)
        run.bold = span.bold or None
        run.italic = span.italic or None
        if span.code:
            run.font.name = _CODE_FONT


def _docx_name(title: str) -> str:
    """OneDrive-safe ``<title>.docx``, without doubling an extension the title carries."""
    return f"{_safe_segment(strip_docx_suffix(title), fallback='document')}.docx"


async def _ensure_folder(client: Any, folder_path: str) -> None:
    """Create each segment of ``folder_path`` under the drive root if it is missing."""
    parent = ""
    for segment in (s for s in folder_path.split("/") if s):
        current = f"{parent}/{segment}" if parent else segment
        get_info = RequestInformation(Method.GET, _ITEM_URL, {"path": current})
        try:
            existing = await client.request_adapter.send_async(get_info, DriveItem, _ERROR_MAP)
        except ODataError as exc:
            # Only a real 404 means "create it" - a 429/5xx must propagate rather
            # than trigger a folder-create write against a folder that may exist.
            if _status_code(exc) != 404:
                raise
            existing = None
        if existing is None:
            new_folder = DriveItem()
            new_folder.name = segment
            new_folder.folder = Folder()
            # Lose the race gracefully: another writer may have just created it.
            new_folder.additional_data = {"@microsoft.graph.conflictBehavior": "replace"}
            if parent:
                post_info = RequestInformation(Method.POST, _CHILDREN_URL, {"path": parent})
            else:
                post_info = RequestInformation(Method.POST, _CHILDREN_ROOT_URL, {})
            post_info.set_content_from_parsable(
                client.request_adapter, "application/json", new_folder
            )
            await client.request_adapter.send_async(post_info, DriveItem, _ERROR_MAP)
        parent = current


async def _get_item_by_id(client: Any, item_id: str) -> DriveItem:
    """``GET /me/drive/items/{id}`` - a 404 becomes ``DriveItemNotFoundError`` (exit 5)."""
    get_info = RequestInformation(Method.GET, _ITEM_BY_ID_URL, {"id": item_id})
    try:
        item = await client.request_adapter.send_async(get_info, DriveItem, _ERROR_MAP)
    except ODataError as exc:
        if _status_code(exc) == 404:
            raise DriveItemNotFoundError(
                f"no document with id {item_id!r} (blumkin can only update docs it created)"
            ) from exc
        raise
    if item is None:
        raise DriveItemNotFoundError(f"no document with id {item_id!r}")
    return item


def _render_block(document: DocxDocument, block: DocBlock, *, ordinal: int) -> None:
    if block.kind == "heading":
        _add_runs(document.add_paragraph(style=f"Heading {min(block.level, 9)}"), block.spans)
        return
    if block.kind == "bullet":
        _add_runs(document.add_paragraph(style=_LIST_STYLES[1 if block.level else 0]), block.spans)
        return
    if block.kind == "number":
        # python-docx's built-in "List Number" style shares one numId, so a
        # second ordered list continues counting instead of restarting. Emit the
        # marker as literal text (`_render_docx` restarts `ordinal` per list) so
        # numbering matches the source, like the Google backend.
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18 * (block.level + 1))
        paragraph.add_run(f"{ordinal}. ")
        _add_runs(paragraph, block.spans)
        return
    if block.kind == "quote":
        # No dedicated Word "block quote" style is used here - a left indent is
        # the closest non-bulleted equivalent, matching the Google backend's
        # `indentStart` (issue #264).
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(36)
        _add_runs(paragraph, block.spans)
        return
    if block.kind == "rule":
        document.add_paragraph("_" * 40)
        return
    if block.kind in ("code", "table"):
        text = block.code_text if block.kind == "code" else table_to_text(block.rows)
        paragraph = document.add_paragraph()
        for offset, line in enumerate(text.split("\n")):
            if offset:
                paragraph.add_run().add_break()
            paragraph.add_run(line).font.name = _CODE_FONT
        return
    _add_runs(document.add_paragraph(), block.spans)


def _render_docx(blocks: list[DocBlock]) -> bytes:
    document = Document()
    counters: dict[int, int] = {}
    previous_kind = ""
    for block in blocks:
        ordinal = 0
        if block.kind == "number":
            if previous_kind != "number":
                counters.clear()
            counters[block.level] = counters.get(block.level, 0) + 1
            for deeper in [level for level in counters if level > block.level]:
                del counters[deeper]
            ordinal = counters[block.level]
        previous_kind = block.kind
        _render_block(document, block, ordinal=ordinal)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _rename_item(client: Any, item_id: str, name: str) -> DriveItem | None:
    """``PATCH /me/drive/items/{id}`` with a new ``name`` (already ``.docx``-suffixed)."""
    patch = DriveItem()
    patch.name = name
    patch_info = RequestInformation(Method.PATCH, _ITEM_BY_ID_URL, {"id": item_id})
    patch_info.set_content_from_parsable(client.request_adapter, "application/json", patch)
    return await client.request_adapter.send_async(patch_info, DriveItem, _ERROR_MAP)


def _safe_folder(folder: str | None) -> str | None:
    if not folder:
        return None
    segments = [_safe_segment(part, fallback="") for part in folder.split("/")]
    kept = [seg for seg in segments if seg]
    return "/".join(kept) or None


def _safe_segment(value: str, *, fallback: str) -> str:
    stripped = "".join(" " if ch in _UNSAFE_SEGMENT_CHARS else ch for ch in value)
    return " ".join(stripped.split()) or fallback


def _status_code(exc: ODataError) -> int | None:
    status = getattr(exc, "response_status_code", None)
    return status if isinstance(status, int) else None
