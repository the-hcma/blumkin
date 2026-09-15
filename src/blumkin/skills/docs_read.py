"""Read local PDF, DOCX, XLSX, and common image files into a stable JSON shape."""

from __future__ import annotations

import importlib
import shutil
from pathlib import Path
from typing import Any

from docx import Document

from blumkin.config import BlumkinConfig
from blumkin.prompt_injection import format_injection_warning_banner, scan_for_injection


class DocsReadExtraMissingError(ValueError):
    """`docs read` needs an optional dependency group that is not installed."""


class DocsReadFileNotFoundError(FileNotFoundError):
    """The requested local file does not exist or is not a regular file."""


class DocsReadOcrUnavailableError(ValueError):
    """OCR is needed (via `--ocr` for PDFs, or always for images), but the
    extra or a required system binary is missing."""


class DocsReadOversizeError(ValueError):
    """The input file or extracted output exceeds the local safety caps."""


class DocsReadUnsupportedFormatError(ValueError):
    """The file extension is not one of the supported local document formats."""


async def docs_read(
    *,
    config: BlumkinConfig,
    path: str,
    pages: str | None = None,
    sheet: str | None = None,
    ocr: bool = False,
) -> dict[str, Any]:
    """Read one local file already on disk; no provider, auth, or network calls."""
    del config
    file_path = Path(path).expanduser()
    _ensure_file(file_path)
    kind = file_path.suffix.lower()
    _validate_flags(kind, ocr=ocr, pages=pages, sheet=sheet)

    ocr_used = False
    if kind == ".docx":
        extracted_pages = _read_docx(file_path)
    elif kind == ".pdf":
        extracted_pages, ocr_used = _read_pdf(file_path, ocr=ocr, pages=pages)
    elif kind == ".xlsx":
        extracted_pages = _read_xlsx(file_path, sheet=sheet)
    elif kind in _IMAGE_EXTENSIONS:
        extracted_pages = _read_image(file_path)
        ocr_used = True
    else:
        supported_images = ", ".join(sorted(_IMAGE_EXTENSIONS))
        raise DocsReadUnsupportedFormatError(
            f"unsupported file type {kind or '<none>'!r} "
            f"(expected .pdf, .docx, .xlsx, or an image: {supported_images})"
        )

    return {
        "injection_warning": _scan_pages_for_injection(extracted_pages),
        "kind": kind.removeprefix("."),
        "ocr_used": ocr_used,
        "ok": True,
        "pages": extracted_pages,
        "path": str(file_path.resolve()),
    }


def format_docs_read_human(payload: dict[str, Any]) -> list[str]:
    pages = payload.get("pages") or []
    lines = [
        f"Read {payload.get('path')!r} ({payload.get('kind')}); "
        f"{len(pages)} section(s); OCR used: {'yes' if payload.get('ocr_used') else 'no'}"
    ]
    lines.extend(format_injection_warning_banner(payload.get("injection_warning")))
    for page in pages:
        heading = f"section {page.get('index')}"
        if page.get("sheet"):
            heading = f"sheet {page.get('sheet')!r}"
        lines.append("")
        lines.append(f"[{heading}]")
        text = str(page.get("text") or "").strip()
        lines.extend(text.splitlines() or ["(no extractable text)"])
        tables = page.get("tables") or []
        if tables:
            lines.append("")
            lines.append(f"tables: {len(tables)}")
            for table_index, table in enumerate(tables, start=1):
                lines.append(f"  table {table_index}:")
                for row in table:
                    lines.append("    " + " | ".join(row))
    return lines


# Pure image formats have no text layer, so OCR is implicit (always on) for
# this kind - unlike PDFs, where `--ocr` is an opt-in fallback for pages that
# have no extractable text. Multi-frame formats (multi-page TIFF, animated
# GIF) are read as a single frame; HEIC/HEIF is not supported yet (needs the
# `pillow-heif` extra).
_IMAGE_EXTENSIONS = frozenset({".bmp", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"})
_MAX_EXTRACTED_BYTES = 1_000_000
_MAX_FILE_BYTES = 25_000_000
_OCR_EXTRA_HINT = "docs read --ocr needs the ocr extra: uv tool install -e '.[pdf,ocr]'"
_POPPLER_HINT = "poppler not found on PATH, install via `brew install tesseract poppler`"
_TESSERACT_HINT = "tesseract not found on PATH, install via `brew install tesseract poppler`"


class _BudgetGuard:
    """Tracks extracted bytes as they are produced and raises as soon as the cap
    is exceeded - checking only after a whole document/sheet is already
    materialized would defeat the cap's purpose (bounding memory)."""

    __slots__ = ("total",)

    def __init__(self) -> None:
        self.total = 0

    def add(self, *chunks: str) -> None:
        self.total += sum(len(chunk.encode("utf-8")) for chunk in chunks)
        if self.total > _MAX_EXTRACTED_BYTES:
            raise DocsReadOversizeError(
                "extracted content exceeds "
                f"{_MAX_EXTRACTED_BYTES} bytes; narrow --pages or choose a smaller file"
            )


def _ensure_file(path: Path) -> None:
    if not path.exists() or not path.is_file():
        raise DocsReadFileNotFoundError(f"file not found: {path}")
    size = path.stat().st_size
    if size > _MAX_FILE_BYTES:
        raise DocsReadOversizeError(
            f"file is larger than {_MAX_FILE_BYTES} bytes; choose a smaller file"
        )


def _import_image_ocr_modules() -> tuple[Any, Any]:
    # Pure image OCR only needs Pillow (to open the file) and pytesseract (to
    # run OCR) - unlike PDF OCR, it never calls pdf2image, so it must not
    # require poppler.
    try:
        pil_image = importlib.import_module("PIL.Image")
        pytesseract = importlib.import_module("pytesseract")
    except ModuleNotFoundError as exc:
        raise DocsReadOcrUnavailableError(_OCR_EXTRA_HINT) from exc
    return pil_image, pytesseract


def _import_ocr_modules() -> tuple[Any, Any, Any]:
    try:
        pdf2image = importlib.import_module("pdf2image")
        pdf2image_exceptions = importlib.import_module("pdf2image.exceptions")
        pytesseract = importlib.import_module("pytesseract")
    except ModuleNotFoundError as exc:
        raise DocsReadOcrUnavailableError(_OCR_EXTRA_HINT) from exc
    return pdf2image, pdf2image_exceptions, pytesseract


def _import_openpyxl() -> Any:
    try:
        return importlib.import_module("openpyxl")
    except ModuleNotFoundError as exc:
        raise DocsReadExtraMissingError(
            "docs read needs the xlsx extra: uv tool install -e '.[xlsx]'"
        ) from exc


def _import_pdfplumber() -> Any:
    try:
        return importlib.import_module("pdfplumber")
    except ModuleNotFoundError as exc:
        raise DocsReadExtraMissingError(
            "docs read needs the pdf extra: uv tool install -e '.[pdf]'"
        ) from exc


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _ocr_pdf_page(path: Path, *, page_index: int) -> str:
    _require_ocr_binaries()
    pdf2image, pdf2image_exceptions, pytesseract = _import_ocr_modules()
    poppler_errors = tuple(
        error
        for name in ("PDFInfoNotInstalledError", "PDFPageCountError", "PopplerNotInstalledError")
        if (error := getattr(pdf2image_exceptions, name, None)) is not None
    )
    try:
        images = pdf2image.convert_from_path(
            str(path),
            first_page=page_index,
            last_page=page_index,
            fmt="png",
            single_file=True,
        )
    except poppler_errors as exc:
        raise DocsReadOcrUnavailableError(_POPPLER_HINT) from exc
    if not images:
        return ""
    try:
        return str(pytesseract.image_to_string(images[0])).strip()
    except getattr(pytesseract, "TesseractNotFoundError", RuntimeError) as exc:
        raise DocsReadOcrUnavailableError(_TESSERACT_HINT) from exc


def _parse_pages(raw: str | None, *, total_pages: int) -> list[int]:
    if raw is None:
        return list(range(1, total_pages + 1))
    selected: set[int] = set()
    chunks = [part.strip() for part in raw.split(",") if part.strip()]
    if not chunks:
        raise ValueError(f"invalid --pages value {raw!r} (expected 1-3 or 1,3,5)")
    for chunk in chunks:
        if "-" in chunk:
            start_text, end_text = [part.strip() for part in chunk.split("-", 1)]
            if not start_text.isdigit() or not end_text.isdigit():
                raise ValueError(f"invalid --pages value {raw!r} (expected 1-3 or 1,3,5)")
            start, end = int(start_text), int(end_text)
            if start < 1 or end < start:
                raise ValueError(f"invalid --pages value {raw!r} (expected 1-3 or 1,3,5)")
            # Validate against total_pages before materializing the range - an
            # unclamped `--pages 1-1000000000` would otherwise build a
            # billion-element set before the out-of-range check ever ran.
            if start > total_pages:
                raise ValueError(
                    f"--pages selects page {start}, but this PDF has {total_pages} page(s)"
                )
            selected.update(range(start, min(end, total_pages) + 1))
            continue
        if not chunk.isdigit() or int(chunk) < 1:
            raise ValueError(f"invalid --pages value {raw!r} (expected 1-3 or 1,3,5)")
        page = int(chunk)
        if page > total_pages:
            raise ValueError(f"--pages selects page {page}, but this PDF has {total_pages} page(s)")
        selected.add(page)
    return sorted(selected)


def _read_docx(path: Path) -> list[dict[str, Any]]:
    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001 - emit a clean usage error message
        raise ValueError(f"cannot read docx file {path}: {exc}") from exc
    budget = _BudgetGuard()
    paragraphs: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            budget.add(text)
            paragraphs.append(text)
    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows = _table_rows(table, budget=budget)
        if rows:
            tables.append(rows)
    page = {"index": 1, "tables": tables, "text": "\n\n".join(paragraphs)}
    return [page]


def _read_image(path: Path) -> list[dict[str, Any]]:
    _require_tesseract_binary()
    pil_image, pytesseract = _import_image_ocr_modules()
    try:
        with pil_image.open(path) as image:
            image.load()
            try:
                text = str(pytesseract.image_to_string(image)).strip()
            except getattr(pytesseract, "TesseractNotFoundError", RuntimeError) as exc:
                raise DocsReadOcrUnavailableError(_TESSERACT_HINT) from exc
    except DocsReadOcrUnavailableError:
        raise
    except Exception as exc:  # noqa: BLE001 - emit a clean usage error message
        raise ValueError(f"cannot read image file {path}: {exc}") from exc
    budget = _BudgetGuard()
    budget.add(text)
    page = {"index": 1, "tables": [], "text": text}
    return [page]


def _read_pdf(path: Path, *, ocr: bool, pages: str | None) -> tuple[list[dict[str, Any]], bool]:
    pdfplumber = _import_pdfplumber()
    budget = _BudgetGuard()
    try:
        with pdfplumber.open(str(path)) as document:
            selected_pages = _parse_pages(pages, total_pages=len(document.pages))
            extracted_pages: list[dict[str, Any]] = []
            ocr_used = False
            for page_index in selected_pages:
                page = document.pages[page_index - 1]
                text = str(page.extract_text() or "").strip()
                if ocr and not text:
                    text = _ocr_pdf_page(path, page_index=page_index)
                    ocr_used = True
                budget.add(text)
                tables = [
                    _table_rows(table, budget=budget) for table in page.extract_tables() or []
                ]
                extracted_page = {"index": page_index, "tables": tables, "text": text}
                extracted_pages.append(extracted_page)
    except DocsReadOcrUnavailableError, DocsReadOversizeError, ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 - emit a clean usage error message
        raise ValueError(f"cannot read pdf file {path}: {exc}") from exc
    return extracted_pages, ocr_used


def _read_xlsx(path: Path, *, sheet: str | None) -> list[dict[str, Any]]:
    openpyxl = _import_openpyxl()
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - emit a clean usage error message
        raise ValueError(f"cannot read xlsx file {path}: {exc}") from exc
    budget = _BudgetGuard()
    try:
        worksheet = _resolve_sheet(workbook, sheet)
        budget.add(worksheet.title)
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            trimmed = _trim_row([_normalize_cell(cell) for cell in row])
            if trimmed:
                # Cell values are returned twice in the payload - once per row in
                # `tables`, once joined into `text` - so budget both renderings,
                # not just the cells, or a wide sheet can double the advertised cap.
                row_text = "\t".join(trimmed)
                budget.add(*trimmed, row_text, "\n")
                rows.append(trimmed)
        page = {
            "index": 1,
            "sheet": worksheet.title,
            "tables": [rows] if rows else [],
            "text": "\n".join("\t".join(row) for row in rows),
        }
    finally:
        workbook.close()
    return [page]


def _require_ocr_binaries() -> None:
    if shutil.which("tesseract") is None:
        raise DocsReadOcrUnavailableError(_TESSERACT_HINT)
    # convert_from_path uses pdfinfo to inspect the PDF and pdftoppm to render it -
    # a partial poppler install (only one of the two) needs both checked, not
    # either, or the missing one surfaces as a generic read error instead.
    if shutil.which("pdfinfo") is None or shutil.which("pdftoppm") is None:
        raise DocsReadOcrUnavailableError(_POPPLER_HINT)


def _require_tesseract_binary() -> None:
    # Pure image OCR never calls pdf2image, so it does not need poppler -
    # only the tesseract binary itself.
    if shutil.which("tesseract") is None:
        raise DocsReadOcrUnavailableError(_TESSERACT_HINT)


def _resolve_sheet(workbook: Any, sheet: str | None) -> Any:
    worksheets = list(workbook.worksheets)
    if not worksheets:
        raise ValueError("workbook has no worksheets")
    if sheet is None:
        return worksheets[0]
    wanted = sheet.strip()
    if not wanted:
        raise ValueError("--sheet must not be blank")
    # Excel allows a numeric worksheet name (e.g. "2024") - check for an exact
    # name match before treating a digit string as a 1-based index.
    if wanted in workbook.sheetnames:
        return workbook[wanted]
    if wanted.isdigit():
        index = int(wanted)
        if 1 <= index <= len(worksheets):
            return worksheets[index - 1]
        raise ValueError(
            f"--sheet index {index} is out of range for {len(worksheets)} worksheet(s)"
        )
    try:
        return workbook[wanted]
    except KeyError as exc:
        raise ValueError(
            f"worksheet {wanted!r} not found; available sheets: {workbook.sheetnames}"
        ) from exc


def _scan_pages_for_injection(pages: list[dict[str, Any]]) -> dict[str, object] | None:
    """Scan every extracted page's text and table cells for prompt-injection patterns.

    Table cells need their own pass because `.docx` extraction keeps table
    text out of `page["text"]` (python-docx's `Document.paragraphs` excludes
    table-cell paragraphs) - a payload placed in a table cell would otherwise
    render under `table N:` with no warning at all.

    Advisory only (see `blumkin.prompt_injection`): a match never blocks or
    alters the extracted content, it only surfaces a warning in the payload
    and the human-formatted output.
    """
    findings = []
    for page in pages:
        index = page.get("index")
        text = str(page.get("text") or "")
        result = scan_for_injection(text, location=f"pages[{index}].text")
        findings.extend(result.findings)
        for table_number, table in enumerate(page.get("tables") or [], start=1):
            table_text = "\n".join("\t".join(row) for row in table)
            result = scan_for_injection(
                table_text, location=f"pages[{index}].tables[{table_number}]"
            )
            findings.extend(result.findings)
    if not findings:
        return None
    return {
        "findings": [
            {"family": f.family, "location": f.location, "snippet": f.snippet} for f in findings
        ],
        "matched": True,
    }


def _table_rows(table: Any, *, budget: _BudgetGuard) -> list[list[str]]:
    rows: list[list[str]] = []
    if hasattr(table, "rows"):
        iterable = ([cell.text for cell in row.cells] for row in table.rows)
    else:
        iterable = table
    for row in iterable:
        trimmed = _trim_row([_normalize_cell(cell) for cell in row])
        if trimmed:
            budget.add(*trimmed)
            rows.append(trimmed)
    return rows


def _trim_row(row: list[str]) -> list[str]:
    trimmed = list(row)
    while trimmed and not trimmed[-1]:
        trimmed.pop()
    return trimmed


def _validate_flags(kind: str, *, ocr: bool, pages: str | None, sheet: str | None) -> None:
    if kind == ".pdf":
        if sheet is not None:
            raise ValueError("--sheet is only valid for .xlsx files")
        return
    if ocr:
        raise ValueError("--ocr is only valid for .pdf files")
    if pages is not None:
        raise ValueError("--pages is only valid for .pdf files")
    if kind in {".docx", *_IMAGE_EXTENSIONS} and sheet is not None:
        raise ValueError("--sheet is only valid for .xlsx files")
